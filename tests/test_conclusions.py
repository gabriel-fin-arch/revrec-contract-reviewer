"""The reviewer's conclusions, read back from a filled-in worklist.

The workbooks here are filled in with openpyxl the way a reviewer would fill
them in Excel: conclusion, name, date in the three right-hand columns. The
runs are rules runs -- no test in this file makes an API call.
"""

from __future__ import annotations

import io
import json
import shutil
from datetime import date
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from revrec_contract_reviewer.assess import assess
from revrec_contract_reviewer.cli import app
from revrec_contract_reviewer.conclusions import (
    CONCLUSIONS_FILE,
    WorklistMismatch,
    collect_conclusions,
    write_conclusions,
)
from revrec_contract_reviewer.extract import ExtractorChoice, extract_contract
from revrec_contract_reviewer.models.document import ContractDocument
from revrec_contract_reviewer.models.review import ContractReview
from revrec_contract_reviewer.report import render_index, render_memo, render_pdf_memo, write_worklist

FILL = ("Closed: the fee is a material right, recognised over the term.", "A. Reviewer", "2026-10-23")


@pytest.fixture(scope="module")
def reviews(documents: dict[str, ContractDocument]) -> list[ContractReview]:
    return [assess(extract_contract(document, ExtractorChoice.RULES)) for document in documents.values()]


def _write_run(reviews: list[ContractReview], run_dir: Path, unreadable=()) -> Path:
    """A run folder: one JSON per contract and the worklist, as `review` leaves them."""
    run_dir.mkdir(parents=True, exist_ok=True)
    for review in reviews:
        (run_dir / f"{review.doc_id}.json").write_text(review.model_dump_json(indent=2), encoding="utf-8")
    return write_worklist(reviews, run_dir, unreadable)


def _fill(worklist: Path, rows: list[int], fill: tuple[str, str, str] = FILL) -> None:
    """Write conclusion, reviewer and date into the given 1-based sheet rows."""
    workbook = load_workbook(worklist)
    items = workbook["open items"]
    for row in rows:
        for column, value in zip((10, 11, 12), fill, strict=True):
            items.cell(row=row, column=column).value = value
    workbook.save(worklist)
    workbook.close()


def _item_rows(worklist: Path) -> list[tuple]:
    workbook = load_workbook(worklist, read_only=True)
    rows = list(workbook["open items"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    return rows


def test_a_filled_in_workbook_is_recorded_beside_the_memos(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    filled = list(range(2, 2 + len(_item_rows(worklist))))
    _fill(worklist, filled)

    read = collect_conclusions(worklist)
    path = write_conclusions(read, tmp_path)

    assert path == tmp_path / CONCLUSIONS_FILE
    assert read.problems == [] and read.open_items == []
    record = json.loads(path.read_text(encoding="utf-8"))
    assert record["run"] == read.run
    assert len(record["conclusions"]) == len(filled)
    for row, written in zip(_item_rows(worklist), record["conclusions"], strict=True):
        # Per contract and item, exactly as typed.
        assert written["contract"] == row[0] and written["item"] == row[3]
        assert written["conclusion"] == FILL[0]
        assert written["reviewer"] == FILL[1] and written["date"] == FILL[2]


def test_a_partially_filled_workbook_leaves_the_rest_open(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    _fill(worklist, [2, 3])

    read = collect_conclusions(worklist)

    assert read.problems == []
    assert len(read.conclusions.conclusions) == 2
    assert len(read.open_items) == len(_item_rows(worklist)) - 2


def test_a_workbook_from_another_run_is_refused(reviews: list[ContractReview], tmp_path: Path):
    one, two = tmp_path / "one", tmp_path / "two"
    worklist = _write_run(reviews, one)
    _write_run(reviews[:-1], two)  # a different run: one contract fewer
    moved = Path(shutil.copy2(worklist, two / worklist.name))  # the first run's workbook beside the second run

    with pytest.raises(WorklistMismatch, match="from run"):
        collect_conclusions(moved)
    assert not (two / CONCLUSIONS_FILE).exists()


def test_a_workbook_whose_run_no_longer_stands_is_refused(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    (tmp_path / f"{reviews[-1].doc_id}.json").unlink()  # the folder no longer holds the run

    with pytest.raises(WorklistMismatch, match="from run"):
        collect_conclusions(worklist)


def test_an_edited_row_is_reported_with_its_row(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    workbook = load_workbook(worklist)
    workbook["open items"]["G2"].value = "evidence the run never wrote"
    workbook.save(worklist)
    workbook.close()
    _fill(worklist, [2, 3])

    read = collect_conclusions(worklist)

    assert read.problems == [
        "row 2: not a row this run raised -- the contract, item or evidence was edited, or the row was added"
    ]
    assert len(read.conclusions.conclusions) == 1


def test_a_row_the_run_did_not_raise_is_reported(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    workbook = load_workbook(worklist)
    added = ["invented-contract", "judgment", "2", "material right", "Is there one?", "", "", "", "", *FILL]
    workbook["open items"].append(added)
    workbook.save(worklist)
    workbook.close()
    row = len(_item_rows(worklist)) + 1  # header plus the rows the run wrote

    read = collect_conclusions(worklist)

    assert read.problems == [
        f"row {row}: not a row this run raised -- the contract, item or evidence was edited, or the row was added"
    ]


def test_a_conclusion_without_a_reviewer_or_a_date_is_reported_not_accepted(
    reviews: list[ContractReview], tmp_path: Path
):
    worklist = _write_run(reviews, tmp_path)
    _fill(worklist, [2], fill=(FILL[0], "", ""))
    _fill(worklist, [3], fill=(FILL[0], FILL[1], ""))

    read = collect_conclusions(worklist)

    assert read.problems == [
        "row 2: a conclusion needs all three of a conclusion, a reviewer and a date",
        "row 3: a conclusion needs all three of a conclusion, a reviewer and a date",
    ]
    assert read.conclusions.conclusions == []
    assert len(read.open_items) == len(_item_rows(worklist))


def test_a_date_excel_turned_into_a_date_comes_back_as_typed(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    workbook = load_workbook(worklist)
    items = workbook["open items"]
    items.cell(row=2, column=10).value = FILL[0]
    items.cell(row=2, column=11).value = FILL[1]
    items.cell(row=2, column=12).value = date(2026, 10, 23)
    workbook.save(worklist)
    workbook.close()

    read = collect_conclusions(worklist)

    assert read.conclusions.conclusions[0].date == "2026-10-23"


def test_a_conclusion_on_an_unreadable_file_is_the_reviewers_own(reviews: list[ContractReview], tmp_path: Path):
    unreadable = [("scanned.pdf", "no text layer: run OCR first")]
    worklist = _write_run(reviews, tmp_path, unreadable)
    _fill(worklist, [2], fill=("Read after OCR: a one-year fixed-fee lease.", FILL[1], FILL[2]))

    read = collect_conclusions(worklist)

    assert read.problems == []
    conclusion = read.conclusions.conclusions[0]
    assert conclusion.contract == "scanned.pdf" and conclusion.kind == "unreadable"
    assert read.unreadable == unreadable


def test_the_memo_index_and_pdf_show_each_item_concluded_or_open(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    rows = _item_rows(worklist)
    judgment_row = next(number for number, row in enumerate(rows, start=2) if row[1] == "judgment")
    _fill(worklist, [judgment_row])
    read = collect_conclusions(worklist)
    conclusions = read.conclusions
    concluded = next(review for review in reviews if review.doc_id == rows[judgment_row - 2][0])

    memo = render_memo(concluded, conclusions=conclusions)
    assert f"Concluded by {FILL[1]} on {FILL[2]}:" in memo
    index = render_index(reviews, conclusions=conclusions)
    assert 'class="tag concluded"' in index and 'class="tag open"' in index

    with pdfplumber.open(io.BytesIO(render_pdf_memo(concluded, conclusions))) as pdf:
        text = "".join(page.extract_text() or "" for page in pdf.pages)
    assert "Concluded by" in text


def test_a_fresh_memo_says_nothing_about_conclusions(reviews: list[ContractReview]):
    memo = render_memo(next(review for review in reviews if review.flags))

    assert "Concluded by" not in memo and "awaiting the reviewer's conclusion" not in memo


def test_conclusions_lookup_matches_flags_and_failed_checks(reviews: list[ContractReview], tmp_path: Path):
    worklist = _write_run(reviews, tmp_path)
    rows = _item_rows(worklist)
    _fill(worklist, list(range(2, 2 + len(rows))))
    conclusions = collect_conclusions(worklist).conclusions

    for review in reviews:
        for flag in review.flags:
            found = conclusions.for_flag(review.doc_id, flag)
            assert found is not None and found.reviewer == FILL[1]
        for check in review.failed_checks():
            assert conclusions.for_check(review.doc_id, check) is not None
    with_flags = next(review for review in reviews if review.flags)
    assert conclusions.for_flag("no-such-contract", with_flags.flags[0]) is None


def test_the_conclusions_command_records_the_fill_and_rewrites_the_memos(tmp_path: Path, monkeypatch):
    """End to end, the way a reviewer runs it: review, fill, conclusions."""
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    runner = CliRunner()
    assert runner.invoke(app, ["generate-corpus"]).exit_code == 0
    assert runner.invoke(app, ["review", "--extractor", "rules"]).exit_code == 0

    worklist = tmp_path / "out" / "worklist.xlsx"
    rows = _item_rows(worklist)
    _fill(worklist, list(range(2, 2 + len(rows))))
    result = runner.invoke(app, ["conclusions", str(worklist)])

    assert result.exit_code == 0, result.output
    record = json.loads((tmp_path / "out" / CONCLUSIONS_FILE).read_text(encoding="utf-8"))
    assert len(record["conclusions"]) == len(rows)
    assert "nothing left open" in result.output
    index = (tmp_path / "out" / "index.html").read_text(encoding="utf-8")
    assert 'class="tag concluded"' in index and 'class="tag open"' not in index


def test_the_conclusions_command_reports_rejected_rows_but_records_the_rest(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    runner = CliRunner()
    assert runner.invoke(app, ["generate-corpus"]).exit_code == 0
    assert runner.invoke(app, ["review", "--extractor", "rules"]).exit_code == 0

    worklist = tmp_path / "out" / "worklist.xlsx"
    _fill(worklist, [2], fill=(FILL[0], "", ""))  # no reviewer, no date: reported, not accepted
    _fill(worklist, [3])

    result = runner.invoke(app, ["conclusions", str(worklist)])

    assert result.exit_code == 1
    assert "row 2" in result.output
    record = json.loads((tmp_path / "out" / CONCLUSIONS_FILE).read_text(encoding="utf-8"))
    assert len(record["conclusions"]) == 1
