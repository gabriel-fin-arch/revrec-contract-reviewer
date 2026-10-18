"""The PDF memo and the worklist workbook.

The PDF is read back with pdfplumber, the same library the pipeline reads
contracts with, and compared with the review it was built from: a memo for
filing that dropped a quote or a page would be worse than no PDF at all.
"""

from __future__ import annotations

import io
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook
from reportlab.lib.units import mm

from revrec_contract_reviewer.assess import assess
from revrec_contract_reviewer.extract import ExtractorChoice, extract_contract
from revrec_contract_reviewer.models.document import ContractDocument
from revrec_contract_reviewer.models.review import ContractReview
from revrec_contract_reviewer.report import render_pdf_memo, write_pdf_memo, write_worklist
from revrec_contract_reviewer.report.memo import _fact_groups, attention_order
from revrec_contract_reviewer.report.worklist import WORKLIST_COLUMNS, worklist_rows


def _review(document: ContractDocument) -> ContractReview:
    return assess(extract_contract(document, ExtractorChoice.RULES))


@pytest.fixture(scope="module")
def reviews(documents: dict[str, ContractDocument]) -> list[ContractReview]:
    return [_review(document) for document in documents.values()]


def _squeezed(text: str) -> str:
    """Whitespace removed: line breaks and word spacing are layout, not content."""
    return "".join(text.split())


# The page footer sits in the bottom 18 mm margin. Reading above it gives the
# memo's own text, which then runs on from one page to the next.
FOOTER = 14 * mm


def _pdf_text(review: ContractReview) -> str:
    with pdfplumber.open(io.BytesIO(render_pdf_memo(review))) as pdf:
        pages = [page.crop((0, 0, page.width, page.height - FOOTER)).extract_text() or "" for page in pdf.pages]
    return _squeezed("\n".join(pages))


def test_every_page_carries_the_contract_and_its_number(helix: ContractDocument):
    with pdfplumber.open(io.BytesIO(render_pdf_memo(_review(helix)))) as pdf:
        footers = [page.crop((0, page.height - FOOTER, page.width, page.height)).extract_text() for page in pdf.pages]

    assert len(footers) > 1
    assert footers == [f"nimbus-of-helix · ASC 606 review · page {n}" for n in range(1, len(footers) + 1)]


def _quoted(citation) -> str:
    return _squeezed(f"{citation.excerpt()} ({citation.label()})")


def test_every_pdf_memo_says_what_its_review_says(reviews: list[ContractReview]):
    """For all twelve contracts, each thing in its place: a judgment's quotes
    inside that judgment, a check next to its result, a value next to its
    label. The same quote often supports a judgment and a schedule row, so
    finding it anywhere in the text would prove nothing about either."""
    for review in reviews:
        text = _pdf_text(review)
        assert text.startswith(_squeezed(review.doc_id))

        for flag in review.flags:
            start = text.index(_squeezed(flag.question))
            block = text[start : text.index(_squeezed(flag.why_a_human), start)]
            assert _squeezed(flag.trigger) in block, (review.doc_id, flag.judgment_type)
            for citation in flag.evidence:
                assert _quoted(citation) in block, (review.doc_id, flag.judgment_type, citation.label())

        for check in review.checks:
            line = f"{check.name}: {check.outcome.value.replace('_', ' ')}. {check.detail}"
            assert _squeezed(line) in text, (review.doc_id, check.name)

        for obligation in review.extraction.obligations:
            assert _squeezed(obligation.label) in text, (review.doc_id, obligation.label)

        schedule = text[text.index(_squeezed("Supporting schedule")) :]
        for group in _fact_groups(review.extraction):
            for row in group.rows:
                assert _squeezed(f"{row.label}: {row.display}") in schedule, (review.doc_id, row.label)
                for citation in row.field.citations[:3]:
                    assert _quoted(citation) in schedule, (review.doc_id, row.label)


def test_a_contract_with_nothing_to_judge_says_so_in_the_pdf(pinegrove: ContractDocument):
    assert _squeezed("Nothing in this document requires a judgment") in _pdf_text(_review(pinegrove))


def test_the_same_review_gives_the_same_pdf(helix: ContractDocument, tmp_path: Path):
    review = _review(helix)

    written = write_pdf_memo(review, tmp_path)

    assert written.name == "nimbus-of-helix.pdf"
    assert written.read_bytes() == render_pdf_memo(review)
    assert written.read_bytes().startswith(b"%PDF-")


def test_every_open_item_in_the_run_is_one_row(reviews: list[ContractReview]):
    rows = worklist_rows(reviews)

    assert len(rows) == sum(len(r.flags) + len(r.failed_checks()) for r in reviews)
    judgments = [dict(zip(WORKLIST_COLUMNS, row, strict=True)) for row in rows if row[1] == "judgment"]
    questions = {flag.question for review in reviews for flag in review.flags}
    assert len(judgments) == sum(len(review.flags) for review in reviews) > 0
    for row in judgments:
        assert row["question or result"] in questions
        assert "p." in row["evidence"]
        assert row["why it needs a person"]
        assert row["reviewer conclusion"] == row["reviewer"] == row["date"] == ""


def test_the_worklist_follows_the_index_order(reviews: list[ContractReview]):
    contracts_in_rows: list[str] = []
    for row in worklist_rows(reviews):
        if row[0] not in contracts_in_rows:
            contracts_in_rows.append(row[0])

    assert contracts_in_rows == [r.doc_id for r in attention_order(reviews) if r.needs_review()]


def test_a_contracts_flags_come_in_the_order_of_the_five_steps(reviews: list[ContractReview]):
    busiest = max(reviews, key=lambda review: len(review.flags))

    steps = [int(row[2]) for row in worklist_rows([busiest]) if row[1] == "judgment"]

    assert steps == sorted(steps)
    assert len(steps) == len(busiest.flags) > 1


def test_files_that_could_not_be_read_come_first(reviews: list[ContractReview]):
    rows = worklist_rows(reviews, [("scanned.pdf", "no text layer: run OCR first")])

    assert rows[0][:5] == ["scanned.pdf", "unreadable", "", "could not be read", "no text layer: run OCR first"]


def test_the_workbook_has_the_items_and_every_contract(reviews: list[ContractReview], tmp_path: Path):
    path = write_worklist(reviews, tmp_path)

    workbook = load_workbook(path, read_only=True)
    items = list(workbook["open items"].iter_rows(values_only=True))
    contracts = list(workbook["contracts"].iter_rows(values_only=True))
    workbook.close()
    assert items[0] == WORKLIST_COLUMNS
    assert len(items) - 1 == len(worklist_rows(reviews))
    assert len(contracts) - 1 == len(reviews)
    clean = {row[0] for row in contracts[1:] if row[3] == "no"}
    assert clean == {r.doc_id for r in reviews if not r.needs_review()}
    assert clean, "the corpus has contracts that need no reviewer, and the sheet should show them"
