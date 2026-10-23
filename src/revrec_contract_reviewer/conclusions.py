"""The reviewer's conclusions, read back from the worklist.

`worklist.xlsx` leaves three columns blank -- reviewer conclusion, reviewer,
date -- for a person to fill in. This module reads them back. It concludes
nothing itself: it records what the reviewer wrote, and says which items are
still open.

The workbook's run sheet names the run it came from (a hash over the run's
reviews, written by `report/worklist.py::run_id`). A workbook whose id doesn't
match the JSON files beside it is from another run, or from a run this folder
no longer holds, and is refused whole. Within a matching workbook every row is
checked against the rows the run raised: a row that doesn't match -- added, or
with its contract, item or evidence edited -- is reported with its row number
and not recorded. So is a conclusion without a reviewer or a date: half a
sign-off is no sign-off.

One gap, stated rather than hidden: rows about files that could not be read
have no JSON behind them, so they can't be matched against the run the way
judgment rows are. A conclusion on one is accepted if the file it names isn't
one of the run's reviewed contracts; the record of it is the reviewer's own.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel

from revrec_contract_reviewer.fileio import write_lf
from revrec_contract_reviewer.models.judgment import JudgmentFlag
from revrec_contract_reviewer.models.review import ContractReview, ReviewCheck
from revrec_contract_reviewer.report.worklist import run_id, worklist_rows

CONCLUSIONS_FILE = "conclusions.json"

# The three right-hand columns of the open items sheet, which the run leaves
# blank and the reviewer fills in.
REVIEWER_COLUMNS = 3


class WorklistMismatch(Exception):
    """The workbook is not one this run can accept conclusions from."""


class Conclusion(BaseModel):
    """One row's reviewer columns, exactly as typed, with the row's identity."""

    contract: str
    kind: str
    step: str
    item: str
    question: str
    conclusion: str
    reviewer: str
    date: str


class Conclusions(BaseModel):
    """What `revrec conclusions` writes: the run id and every accepted row."""

    run: str
    conclusions: list[Conclusion] = []

    def _by_key(self) -> dict[tuple[str, str, str, str], Conclusion]:
        return {(c.contract, c.kind, c.item, c.question): c for c in self.conclusions}

    def for_flag(self, doc_id: str, flag: JudgmentFlag) -> Conclusion | None:
        """The conclusion for one judgment, if the reviewer recorded one."""
        key = (doc_id, "judgment", flag.judgment_type.value.replace("_", " "), flag.question)
        return self._by_key().get(key)

    def for_check(self, doc_id: str, check: ReviewCheck) -> Conclusion | None:
        """The conclusion for one failed check, if the reviewer recorded one."""
        return self._by_key().get((doc_id, "failed check", check.name, check.detail))


@dataclass
class ReadBack:
    """The result of reading a filled-in workbook against its run."""

    run: str
    conclusions: Conclusions
    open_items: list[str] = field(default_factory=list)  # "contract -- item", for the summary
    problems: list[str] = field(default_factory=list)  # rejected rows, each naming its row
    reviews: list[ContractReview] = field(default_factory=list)
    unreadable: list[tuple[str, str]] = field(default_factory=list)


def _typed(value: Any) -> str:
    """A cell as text. Excel turns some dates into datetimes; those go back to ISO."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _load_run(run_dir: Path) -> list[ContractReview]:
    """The run beside the workbook: one JSON per reviewed contract."""
    return [
        ContractReview.model_validate_json(path.read_text(encoding="utf-8"))
        for path in sorted(run_dir.glob("*.json"))
        if path.name != CONCLUSIONS_FILE
    ]


def collect_conclusions(worklist_path: Path) -> ReadBack:
    """Read a filled-in worklist against the run beside it.

    Raises WorklistMismatch when the workbook isn't one this tool wrote, or
    belongs to another run. Row-level problems -- a row the run didn't raise,
    an edited row, a conclusion without a reviewer or a date -- don't stop the
    read: they're reported in the result and the sound rows are still recorded.
    """
    workbook = load_workbook(worklist_path)
    if "open items" not in workbook.sheetnames or "run" not in workbook.sheetnames:
        raise WorklistMismatch(f"{worklist_path.name} is not a worklist written by this tool")
    properties = {
        _typed(key): _typed(value)
        for key, value in workbook["run"].iter_rows(min_col=1, max_col=2, values_only=True)
        if key
    }
    workbook_run = properties.get("run", "")

    reviews = _load_run(worklist_path.parent)
    if not reviews:
        raise WorklistMismatch(f"no reviewed contracts beside {worklist_path.name}: the run's JSON files are missing")
    if workbook_run != run_id(reviews):
        raise WorklistMismatch(
            f"this worklist is from run {workbook_run or '(none named)'}, "
            f"but the folder beside it holds run {run_id(reviews)} -- conclusions go to the run's own worklist"
        )

    # The rows the run raised, keyed by everything but the reviewer columns,
    # so an edit to the contract, the item or the evidence breaks the match.
    expected: dict[tuple[str, ...], None] = {}
    for row in worklist_rows(reviews):
        expected[tuple(row[: len(row) - REVIEWER_COLUMNS])] = None
    reviewed = {review.doc_id for review in reviews}

    read = ReadBack(run=workbook_run, conclusions=Conclusions(run=workbook_run), reviews=reviews)
    concluded: dict[tuple[str, ...], int] = {}
    items = workbook["open items"]
    for number, sheet_row in enumerate(items.iter_rows(min_row=2, values_only=True), start=2):
        cells = [_typed(value) for value in sheet_row]
        if not any(cells):
            continue
        key = tuple(cells[: len(cells) - REVIEWER_COLUMNS])
        conclusion, reviewer, when = cells[-REVIEWER_COLUMNS:]

        if key[1] == "unreadable":
            # No JSON behind these rows to match against; see the docstring.
            read.unreadable.append((key[0], key[4]))
            valid = key[0].removesuffix(".pdf") not in reviewed and key[3] == "could not be read"
        else:
            valid = key in expected
        if not valid:
            read.problems.append(
                f"row {number}: not a row this run raised -- the contract, item or evidence "
                "was edited, or the row was added"
            )
            continue

        if not any((conclusion, reviewer, when)):
            read.open_items.append(f"{key[0]} -- {key[3]}")
            continue
        if not all((conclusion, reviewer, when)):
            read.problems.append(
                f"row {number}: a conclusion needs all three of a conclusion, a reviewer and a date"
            )
            # Not accepted -- and not concluded either, so the item stays open.
            read.open_items.append(f"{key[0]} -- {key[3]}")
            continue
        if key in concluded:
            read.problems.append(f"row {number}: already concluded on row {concluded[key]}")
            continue
        concluded[key] = number
        read.conclusions.conclusions.append(
            Conclusion(
                contract=key[0],
                kind=key[1],
                step=key[2],
                item=key[3],
                question=key[4],
                conclusion=conclusion,
                reviewer=reviewer,
                date=when,
            )
        )
    workbook.close()
    return read


def write_conclusions(read: ReadBack, out_dir: Path) -> Path:
    """Write conclusions.json beside the run's memos and return its path."""
    return write_lf(out_dir / CONCLUSIONS_FILE, read.conclusions.model_dump_json(indent=2) + "\n")


def load_conclusions(run_dir: Path) -> Conclusions | None:
    """The recorded conclusions for a run, if `revrec conclusions` has run there."""
    path = run_dir / CONCLUSIONS_FILE
    if not path.is_file():
        return None
    return Conclusions.model_validate(json.loads(path.read_text(encoding="utf-8")))
