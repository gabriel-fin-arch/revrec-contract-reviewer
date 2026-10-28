"""recalc/rr_recalc.xlsx, the accountant's own reading of three contracts, reconciled against the tool.

recalc/ holds Gabriel's independent recalculation in Excel (#57): for each
contract, the fields the tool reads and the price of each performance
obligation, read by hand from the PDF, in the `recalc` sheet
(doc_id | field | my_value | page | clause). Nothing here writes to it. This
reviews the same PDFs with the rules extractor, the one `revrec review
--extractor rules` uses, which never calls out, and compares row by row:

- every field in the sheet against the tool's value for it; an empty cell and
  a field the tool marks absent are the same answer
- every obligation kind and its price, both ways: a kind on one side only is a
  difference too, and the same price under two different kinds is one

Neither side is adjusted to make the two agree. A difference is looked at by
hand; if it stays, it is recorded in evals/recalc_known_differences.toml with
its issue, and that is Gabriel's decision, not a fix. What a held-out contract
shows is never a reason to change the rules or the prompt (#57). The check
fails on any difference that isn't recorded there, and on any recorded one
that no longer occurs, so the list can't outlive what it explains. CI runs it
on every push and pull request; locally,

    uv run python -m revrec_contract_reviewer.recalc_check
"""

from __future__ import annotations

import sys
import tomllib
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from revrec_contract_reviewer.extract import ExtractorChoice
from revrec_contract_reviewer.models.extraction import ContractExtraction
from revrec_contract_reviewer.models.fields import ExtractedField
from revrec_contract_reviewer.pipeline import review_contract

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "recalc" / "rr_recalc.xlsx"
KNOWN = ROOT / "evals" / "recalc_known_differences.toml"
SHEET = "recalc"
# The development set, then the held-out test set (cli.SET_DIRS).
CORPUS_DIRS = (ROOT / "corpus", ROOT / "corpus" / "test")
OBLIGATION = "obligation:"
NO_ROW = "(no row)"

FIELDS: dict[str, Callable[[ContractExtraction], ExtractedField[Any]]] = {
    "customer": lambda e: e.customer,
    "supplier": lambda e: e.supplier,
    "agreement_type": lambda e: e.agreement_type,
    "currency": lambda e: e.currency,
    "effective_date": lambda e: e.term.effective_date,
    "stated_term_months": lambda e: e.term.stated_term_months,
    "termination_for_convenience": lambda e: e.term.termination_for_convenience,
    "termination_notice_days": lambda e: e.term.termination_notice_days,
    "net_days": lambda e: e.payment.net_days,
    "total_fixed_consideration": lambda e: e.total_fixed_consideration,
}


@dataclass(frozen=True)
class Row:
    doc_id: str
    field: str
    value: object


@dataclass(frozen=True)
class Difference:
    doc_id: str
    field: str
    recalc: object
    tool: object

    def key(self) -> tuple[str, str, str, str]:
        return (self.doc_id, self.field, _show(self.recalc), _show(self.tool))


@dataclass(frozen=True)
class Known:
    """A difference looked at and recorded, with its issue; the values as the report shows them."""

    doc_id: str
    field: str
    recalc: str
    tool: str
    issue: int

    def key(self) -> tuple[str, str, str, str]:
        return (self.doc_id, self.field, self.recalc, self.tool)


@dataclass(frozen=True)
class Verdict:
    recorded: list[tuple[Difference, Known]]
    unrecorded: list[Difference]
    stale: list[Known]

    @property
    def passes(self) -> bool:
        return not self.unrecorded and not self.stale


def read_sheet(path: Path = WORKBOOK) -> list[Row]:
    """The `recalc` sheet's rows, with the values Excel last calculated and saved."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[SHEET].iter_rows(values_only=True)
        header = next(rows)
        if tuple(header[:3]) != ("doc_id", "field", "my_value"):
            raise ValueError(f"{path.name}: the {SHEET!r} sheet should start doc_id | field | my_value, not {header}")
        return [Row(str(doc_id), str(field), value) for doc_id, field, value, *_ in rows if doc_id]
    finally:
        workbook.close()


def find_pdf(doc_id: str) -> Path | None:
    return next((path for d in CORPUS_DIRS if (path := d / f"{doc_id}.pdf").is_file()), None)


def tool_extractions(doc_ids: set[str]) -> dict[str, ContractExtraction | None]:
    """What the rules extractor reads from each contract; None for a doc_id with no PDF in the corpus."""
    found: dict[str, ContractExtraction | None] = {}
    for doc_id in sorted(doc_ids):
        pdf = find_pdf(doc_id)
        found[doc_id] = review_contract(pdf, ExtractorChoice.RULES).extraction if pdf else None
    return found


def obligation_prices(extraction: ContractExtraction) -> dict[str, Decimal | None]:
    """Stated price per obligation kind, summed if a kind appears twice; None if any of its prices is absent."""
    prices: dict[str, Decimal | None] = {}
    for obligation in extraction.obligations:
        kind, price = obligation.kind.value, obligation.stated_price.value
        if kind not in prices:
            prices[kind] = price
        else:
            earlier = prices[kind]
            prices[kind] = None if earlier is None or price is None else earlier + price
    return prices


def _normal(value: object) -> tuple[str, object]:
    """One comparable form per kind of value, so a number never equals a bool and '' equals an absent field."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return ("empty", None)
    if isinstance(value, bool):
        return ("bool", value)
    if isinstance(value, int | float | Decimal):
        return ("number", Decimal(str(value)))
    if isinstance(value, datetime):
        return ("date", value.date().isoformat())
    if isinstance(value, date):
        return ("date", value.isoformat())
    text = str(value).strip()
    return ("date", text) if _looks_like_date(text) else ("text", text)


def _looks_like_date(text: str) -> bool:
    try:
        date.fromisoformat(text)
    except ValueError:
        return False
    return True


def _show(value: object) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "(empty)"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | Decimal):
        return f"{value:,}"
    return str(value)


def differences(rows: list[Row], extractions: dict[str, ContractExtraction | None]) -> list[Difference]:
    """Every row where the sheet and the tool disagree, and every obligation kind only one of them has."""
    found: list[Difference] = []
    by_doc: dict[str, list[Row]] = {}
    for row in rows:
        by_doc.setdefault(row.doc_id, []).append(row)
    for doc_id, doc_rows in by_doc.items():
        extraction = extractions.get(doc_id)
        if extraction is None:
            found.append(Difference(doc_id, "(contract)", "in the sheet", "no PDF in corpus/ or corpus/test/"))
            continue
        in_sheet: dict[str, object] = {}
        for row in doc_rows:
            if row.field.startswith(OBLIGATION):
                in_sheet[row.field.removeprefix(OBLIGATION)] = row.value
            elif row.field not in FIELDS:
                found.append(Difference(doc_id, row.field, row.value, "(not a field this check knows)"))
            elif _normal(row.value) != _normal(tool := FIELDS[row.field](extraction).value):
                found.append(Difference(doc_id, row.field, row.value, tool))
        found += _obligation_differences(doc_id, in_sheet, obligation_prices(extraction))
    return found


def _obligation_differences(
    doc_id: str, in_sheet: dict[str, object], prices: dict[str, Decimal | None]
) -> list[Difference]:
    found = [
        Difference(doc_id, OBLIGATION + kind, mine, prices[kind])
        for kind, mine in in_sheet.items()
        if kind in prices and _normal(mine) != _normal(prices[kind])
    ]
    only_tool = [kind for kind in prices if kind not in in_sheet]
    for kind, mine in in_sheet.items():
        if kind in prices:
            continue
        # The same price under another kind is one difference, about the kind.
        other = next((k for k in only_tool if _normal(prices[k]) == _normal(mine)), None)
        if other is None:
            found.append(Difference(doc_id, OBLIGATION + kind, mine, NO_ROW))
        else:
            only_tool.remove(other)
            found.append(Difference(doc_id, f"obligation kind, price {_show(mine)}", kind, other))
    found += [Difference(doc_id, OBLIGATION + kind, NO_ROW, prices[kind]) for kind in only_tool]
    return found


def load_known(path: Path = KNOWN) -> list[Known]:
    """The recorded differences; none if the file isn't there."""
    if not path.is_file():
        return []
    entries = tomllib.loads(path.read_text(encoding="utf-8")).get("difference", [])
    known: list[Known] = []
    for number, entry in enumerate(entries, start=1):
        missing = {"doc_id", "field", "recalc", "tool", "issue"} - entry.keys()
        if missing:
            raise ValueError(f"{path.name}, difference {number}: no {', '.join(sorted(missing))}")
        if not isinstance(entry["issue"], int) or entry["issue"] < 1:
            raise ValueError(f"{path.name}, difference {number}: issue must be the issue's number")
        known.append(
            Known(str(entry["doc_id"]), str(entry["field"]), str(entry["recalc"]), str(entry["tool"]), entry["issue"])
        )
    return known


def judge(found: list[Difference], known: list[Known]) -> Verdict:
    by_key = {k.key(): k for k in known}
    found_keys = {d.key() for d in found}
    return Verdict(
        recorded=[(d, by_key[d.key()]) for d in found if d.key() in by_key],
        unrecorded=[d for d in found if d.key() not in by_key],
        stale=[k for k in known if k.key() not in found_keys],
    )


def report(rows: list[Row], verdict: Verdict) -> list[str]:
    """Markdown: rows compared per contract, then every difference and whether it's recorded."""
    docs = sorted({row.doc_id for row in rows})
    found = [d for d, _ in verdict.recorded] + verdict.unrecorded
    lines = [
        "## recalc/rr_recalc.xlsx against the tool (rules extractor)",
        "",
        "| Contract | Rows compared | Differences | Recorded with an issue |",
        "|---|---:|---:|---:|",
        *(
            f"| {doc} | {sum(r.doc_id == doc for r in rows)} | {sum(d.doc_id == doc for d in found)} "
            f"| {sum(d.doc_id == doc for d, _ in verdict.recorded)} |"
            for doc in docs
        ),
        "",
    ]
    if found:
        issue = {d.key(): f"#{k.issue}" for d, k in verdict.recorded}
        lines += [
            "| Contract | Field | recalc/rr_recalc.xlsx | tool | Recorded |",
            "|---|---|---|---|---|",
            *(
                f"| {d.doc_id} | {d.field} | {_show(d.recalc)} | {_show(d.tool)} | {issue.get(d.key(), '**no**')} |"
                for d in found
            ),
            "",
        ]
    if verdict.stale:
        lines += [
            f"Recorded in {KNOWN.relative_to(ROOT).as_posix()} but no longer found; take each out once its "
            "issue is settled:",
            "",
            *(f"- {k.doc_id}, {k.field}: {k.recalc} vs {k.tool} (#{k.issue})" for k in verdict.stale),
            "",
        ]
    if verdict.passes:
        recorded = f"each of the {len(found)} difference(s) is recorded with its issue." if found else "no differences."
        lines.append(f"Passes: {len(rows)} rows of {len(docs)} contracts compared one by one; {recorded}")
    else:
        lines.append(
            f"**Fails:** {len(verdict.unrecorded)} difference(s) not recorded, {len(verdict.stale)} recorded "
            "one(s) no longer found. Look at each by hand; neither side gets adjusted to make this pass."
        )
    return lines


def main() -> None:
    rows = read_sheet()
    verdict = judge(differences(rows, tool_extractions({row.doc_id for row in rows})), load_known())
    print("\n".join(report(rows, verdict)))
    if not verdict.passes:
        sys.exit(1)


if __name__ == "__main__":
    main()
